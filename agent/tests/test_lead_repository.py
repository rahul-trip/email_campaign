"""agent/tests/test_lead_repository.py — Tests for XlsxLeadRepository."""

from __future__ import annotations

import openpyxl

from agent.domain.models import COLUMN_ORDER, Lead, VerifiedEmail
from agent.persistence.lead_repository import XlsxLeadRepository
from .conftest import SEED_COMPANIES, SEED_CONTACT


def test_existing_domains_contains_seeded_row(repo: XlsxLeadRepository):
    """existing_domains() must return the domain from the pre-seeded row."""
    domains = repo.existing_domains()
    assert "existingco.com" in domains


def test_row_count_after_seed(repo: XlsxLeadRepository):
    """row_count() should return 2 (header + 1 seed row) on a fresh fixture."""
    assert repo.row_count() == 2


def test_append_leads_increases_row_count(repo: XlsxLeadRepository, sample_lead: Lead):
    """Appending N leads should grow row_count by N."""
    before = repo.row_count()
    n = repo.append_leads([sample_lead, sample_lead])
    assert n == 2
    assert repo.row_count() == before + 2


def test_append_leads_writes_correct_columns(repo: XlsxLeadRepository, sample_lead: Lead, tmp_xlsx):
    """After appending, newly-written columns must match Lead.to_row()."""
    repo.append_leads([sample_lead])

    wb = openpyxl.load_workbook(str(tmp_xlsx), read_only=True, data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # Row 0 = header, row 1 = seed, row 2 = appended
    new_row = all_rows[2]
    header = [str(c).strip() if c else "" for c in all_rows[0]]
    row_dict = {header[i]: new_row[i] for i in range(len(header)) if i < len(new_row)}

    expected = sample_lead.to_row()
    for col, val in expected.items():
        assert row_dict.get(col) == val, f"Column {col!r}: expected {val!r}, got {row_dict.get(col)!r}"


def test_existing_rows_untouched(repo: XlsxLeadRepository, sample_lead: Lead, tmp_xlsx):
    """Appending must not modify any pre-existing row."""
    wb_before = openpyxl.load_workbook(str(tmp_xlsx), read_only=True, data_only=True)
    rows_before = list(wb_before.active.iter_rows(values_only=True))
    wb_before.close()

    repo.append_leads([sample_lead])

    wb_after = openpyxl.load_workbook(str(tmp_xlsx), read_only=True, data_only=True)
    rows_after = list(wb_after.active.iter_rows(values_only=True))
    wb_after.close()

    # Header and seed row should be identical
    assert rows_before[0] == rows_after[0], "Header row was modified"
    assert rows_before[1] == rows_after[1], "Seed row was modified"


def test_existing_domains_reflects_appended(repo: XlsxLeadRepository, sample_lead: Lead):
    """Domains of appended rows must appear in existing_domains() on next call."""
    repo.append_leads([sample_lead])
    domains = repo.existing_domains()
    assert "renewco.io" in domains


def test_append_empty_list(repo: XlsxLeadRepository):
    """Appending an empty list should return 0 and not touch the workbook."""
    before = repo.row_count()
    n = repo.append_leads([])
    assert n == 0
    assert repo.row_count() == before

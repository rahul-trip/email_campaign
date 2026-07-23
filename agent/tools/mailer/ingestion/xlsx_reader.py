"""Reads exactly the three columns this workflow needs from the campaign matrix.

Every other column (Follow-up Text, Points of Commonality, POC Details &
Eccentricities, Specific Person to Contact) is ignored on purpose — that context
is already baked into "Tailored Pitch Email" by the upstream copywriting step.
"""

from __future__ import annotations

from dataclasses import dataclass

import openpyxl

REQUIRED_COLUMNS = ("Direct Business Email", "Organization URL", "Tailored Pitch Email")


@dataclass(frozen=True)
class CampaignRow:
    row_number: int  # 1-indexed spreadsheet row, for error reporting
    to_addr: str
    subject: str
    body: str


class XlsxSchemaError(ValueError):
    """Raised when the workbook is missing one of the required columns."""


def _clean_subject(raw: str) -> str:
    # Organization URL cells contain "Company Name\nhttps://...". Header values
    # can't contain raw newlines, so collapse to a single line.
    return " ".join(raw.split())


def read_campaign_rows(
    xlsx_path: str,
    start_row: int | None = None,
) -> list[CampaignRow]:
    """Read campaign rows from the workbook.

    Args:
        xlsx_path: Path to the campaign matrix workbook.
        start_row: When set, skip rows whose spreadsheet row number is less
                   than this value (1-indexed, header is row 1).  Use this
                   after the agent appends new rows so only the new rows are
                   drafted.  Default None = read all rows (existing behavior).
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    header = [str(cell).strip() if cell is not None else "" for cell in next(ws.iter_rows(max_row=1, values_only=True))]
    col_index = {name: header.index(name) for name in REQUIRED_COLUMNS if name in header}

    missing = [name for name in REQUIRED_COLUMNS if name not in col_index]
    if missing:
        raise XlsxSchemaError(f"Workbook is missing required column(s): {missing}")

    rows: list[CampaignRow] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # start_row filter: skip rows that were already drafted in a prior run
        if start_row is not None and row_number < start_row:
            continue

        to_addr = values[col_index["Direct Business Email"]]
        org_url = values[col_index["Organization URL"]]
        pitch = values[col_index["Tailored Pitch Email"]]

        to_addr = (to_addr or "").strip() if isinstance(to_addr, str) else to_addr
        pitch = (pitch or "").strip() if isinstance(pitch, str) else pitch
        org_url = (org_url or "").strip() if isinstance(org_url, str) else org_url

        if not to_addr or not pitch:
            # Skip rows with no destination address or no drafted body — nothing
            # meaningful to save as a draft.
            continue

        rows.append(
            CampaignRow(
                row_number=row_number,
                to_addr=to_addr,
                subject=_clean_subject(org_url or ""),
                body=pitch,
            )
        )

    return rows
